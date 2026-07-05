"""Word (.docx) dosyasından düz metin çıkarma.

Paragraflar ve tablo hücreleri belge içindeki sırasıyla, satır sonlarıyla
birleştirilir.
"""

from __future__ import annotations

import logging
import re
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

logger = logging.getLogger(__name__)


class DocxError(RuntimeError):
    """Word dosyası okunamadı."""


def _table_text(table: Table) -> list[str]:
    lines: list[str] = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        line = " | ".join(c for c in cells if c)
        if line:
            lines.append(line)
    return lines


def extract_text(docx_path: Path) -> str:
    """Belgedeki tüm metni (paragraflar + tablolar, belge sırasıyla) döndürür."""
    try:
        document = Document(str(docx_path))
    except Exception as exc:  # noqa: BLE001 — python-docx çeşitli hatalar fırlatabilir
        raise DocxError(f"Word dosyası açılamadı: {exc}") from exc

    parts: list[str] = []
    try:
        # python-docx >= 1.1: paragraf ve tabloları belge sırasıyla verir.
        for item in document.iter_inner_content():
            if isinstance(item, Paragraph):
                parts.append(item.text)
            elif isinstance(item, Table):
                parts.extend(_table_text(item))
    except AttributeError:
        # Eski python-docx: önce paragraflar, sonra tablolar.
        parts.extend(p.text for p in document.paragraphs)
        for table in document.tables:
            parts.extend(_table_text(table))

    text = "\n".join(parts)
    # 3+ ardışık boş satırı teke indir, baş/son boşlukları kırp.
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    logger.info("Word'den %d karakter metin çıkarıldı: %s", len(text), docx_path.name)
    return text
